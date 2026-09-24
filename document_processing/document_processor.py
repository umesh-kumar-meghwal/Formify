"""
Dynamic Document Processing + RAG Context Pipeline

SIH 2026 - Member 4: Document Processing + RAG

Supports:
PDF, TXT, DOCX and common image files.

PDFs with selectable text are processed normally.
Scanned/image-only PDFs automatically use OCR.

This file does NOT call Gemini.
Member 1 will use the output of process_document()
for the LLM / Gemini layer.
"""

from pathlib import Path
import re


STOPWORDS = {
    "a", "an", "the", "and", "or", "but", "if", "then", "than",
    "to", "of", "in", "on", "at", "by", "for", "from", "with",
    "about", "into", "over", "under", "after", "before", "during",
    "is", "are", "was", "were", "be", "been", "being",
    "this", "that", "these", "those", "it", "its",
    "as", "what", "which", "who", "where", "when", "why",
    "how", "can", "could", "should", "would", "do", "does",
    "did", "has", "have", "had", "will", "may", "might"
}


SUPPORTED_EXTENSIONS = {
    ".pdf": "PDF",
    ".txt": "TXT",
    ".docx": "DOCX",
    ".png": "IMAGE",
    ".jpg": "IMAGE",
    ".jpeg": "IMAGE",
    ".bmp": "IMAGE",
    ".tiff": "IMAGE",
    ".webp": "IMAGE"
}


# ---------------------------------------------------------
# TEXT CLEANING
# ---------------------------------------------------------

def clean_text(text):
    """Clean extracted text."""

    if not text:
        return ""

    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)

    return text.strip()


# ---------------------------------------------------------
# TOKENIZATION
# ---------------------------------------------------------

def tokenize(text):
    """Convert text into words."""

    return re.findall(
        r"\b[a-zA-Z0-9][a-zA-Z0-9_-]*\b",
        text.lower()
    )


def meaningful_tokens(text):
    """Remove common words."""

    return {
        word
        for word in tokenize(text)
        if word not in STOPWORDS and len(word) > 1
    }


# ---------------------------------------------------------
# CHUNKING
# ---------------------------------------------------------

def chunk_text(
    text,
    source,
    page_number,
    start_chunk_id,
    chunk_size=500,
    overlap=80
):
    """Split text into overlapping chunks."""

    words = text.split()

    if not words:
        return []

    if chunk_size <= 0:
        raise ValueError("chunk_size must be greater than 0.")

    if overlap < 0 or overlap >= chunk_size:
        raise ValueError(
            "overlap must be smaller than chunk_size."
        )

    chunks = []

    step = chunk_size - overlap

    for start in range(0, len(words), step):

        piece = " ".join(
            words[start:start + chunk_size]
        ).strip()

        if not piece:
            continue

        chunks.append({
            "chunk_id": start_chunk_id + len(chunks),
            "source": source,
            "page": page_number,
            "text": piece
        })

        if start + chunk_size >= len(words):
            break

    return chunks


# ---------------------------------------------------------
# TESSERACT SETUP
# ---------------------------------------------------------

def configure_tesseract():
    """
    Configure Tesseract OCR.

    Tesseract is installed at the default Windows location.
    """

    try:
        import pytesseract
    except ImportError:
        raise ImportError(
            "pytesseract is required for OCR. "
            "Run: pip install pytesseract"
        )

    tesseract_path = Path(
        r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    )

    if not tesseract_path.exists():
        raise FileNotFoundError(
            "Tesseract OCR was not found at: "
            f"{tesseract_path}"
        )

    pytesseract.pytesseract.tesseract_cmd = str(
        tesseract_path
    )

    return pytesseract


# ---------------------------------------------------------
# OCR FOR PDF PAGE
# ---------------------------------------------------------

def ocr_pdf_page(page):
    """
    Convert one PDF page into an image
    and extract text using Tesseract OCR.
    """

    try:
        import pymupdf
    except ImportError:
        raise ImportError(
            "PyMuPDF is required. "
            "Run: pip install pymupdf"
        )

    pytesseract = configure_tesseract()

    # Render PDF page as image
    pix = page.get_pixmap(
        matrix=pymupdf.Matrix(2, 2),
        alpha=False
    )

    image_bytes = pix.tobytes("png")

    from PIL import Image
    from io import BytesIO

    image = Image.open(
        BytesIO(image_bytes)
    )

    text = pytesseract.image_to_string(
        image
    )

    return clean_text(text)


# ---------------------------------------------------------
# PDF EXTRACTION + OCR FALLBACK
# ---------------------------------------------------------

def extract_from_pdf(file_path):
    """
    Extract text from PDF page by page.

    Normal PDFs:
        PDF text extraction is used.

    Scanned PDFs:
        If a page has no selectable text,
        OCR is automatically used for that page.
    """

    try:
        import pymupdf
    except ImportError:
        raise ImportError(
            "PyMuPDF is required. "
            "Run: pip install pymupdf"
        )

    doc = pymupdf.open(file_path)

    pages = []
    full_text = []

    try:

        for page_number, page in enumerate(
            doc,
            start=1
        ):

            # First try normal text extraction
            text = page.get_text("text")
            text = clean_text(text)

            # If page has no text, use OCR
            if not text:

                print(
                    f"OCR used for PDF page {page_number}"
                )

                text = ocr_pdf_page(page)

            if text:

                pages.append({
                    "page": page_number,
                    "text": text
                })

                full_text.append(text)

    finally:
        doc.close()

    return "\n\n".join(full_text), pages


# ---------------------------------------------------------
# TXT EXTRACTION
# ---------------------------------------------------------

def extract_from_txt(file_path):
    """Extract text from TXT."""

    path = Path(file_path)

    try:
        text = path.read_text(
            encoding="utf-8-sig"
        )

    except UnicodeDecodeError:

        text = path.read_text(
            encoding="utf-8",
            errors="replace"
        )

    text = clean_text(text)

    pages = []

    if text:

        pages.append({
            "page": None,
            "text": text
        })

    return text, pages


# ---------------------------------------------------------
# DOCX EXTRACTION
# ---------------------------------------------------------

def extract_from_docx(file_path):
    """Extract paragraphs and tables from DOCX."""

    try:
        from docx import Document

    except ImportError:

        raise ImportError(
            "python-docx is required. "
            "Run: pip install python-docx"
        )

    document = Document(file_path)

    parts = []

    # Paragraphs
    for paragraph in document.paragraphs:

        text = clean_text(
            paragraph.text
        )

        if text:
            parts.append(text)

    # Tables
    for table in document.tables:

        for row in table.rows:

            cells = []

            for cell in row.cells:

                text = clean_text(
                    cell.text
                )

                if text:
                    cells.append(text)

            if cells:
                parts.append(
                    " | ".join(cells)
                )

    text = clean_text(
        "\n".join(parts)
    )

    pages = []

    if text:

        pages.append({
            "page": None,
            "text": text
        })

    return text, pages


# ---------------------------------------------------------
# IMAGE OCR
# ---------------------------------------------------------

def extract_from_image(file_path):
    """Extract text from image using OCR."""

    try:
        from PIL import Image
        import pytesseract

    except ImportError:

        raise ImportError(
            "Image OCR requires Pillow and pytesseract. "
            "Run: pip install pillow pytesseract"
        )

    configure_tesseract()

    image = Image.open(file_path)

    text = pytesseract.image_to_string(
        image
    )

    text = clean_text(text)

    pages = []

    if text:

        pages.append({
            "page": None,
            "text": text
        })

    return text, pages


# ---------------------------------------------------------
# DOCUMENT EXTRACTION
# ---------------------------------------------------------

def extract_document(file_path):
    """
    Automatically detect file type and extract text.
    """

    path = Path(file_path)

    if not path.exists():

        raise FileNotFoundError(
            f"Document not found: {file_path}"
        )

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:

        supported = ", ".join(
            SUPPORTED_EXTENSIONS.keys()
        )

        raise ValueError(
            f"Unsupported file type: {extension}. "
            f"Supported: {supported}"
        )

    file_type = SUPPORTED_EXTENSIONS[
        extension
    ]

    if file_type == "PDF":

        text, pages = extract_from_pdf(
            file_path
        )

    elif file_type == "TXT":

        text, pages = extract_from_txt(
            file_path
        )

    elif file_type == "DOCX":

        text, pages = extract_from_docx(
            file_path
        )

    elif file_type == "IMAGE":

        text, pages = extract_from_image(
            file_path
        )

    else:

        raise ValueError(
            "Unsupported file type."
        )

    if not text.strip():

        raise ValueError(
            f"No readable text found in {path.name}"
        )

    return {
        "source": path.name,
        "file_type": file_type,
        "text": text,
        "pages": pages,
        "characters": len(text)
    }


# ---------------------------------------------------------
# BUILD CHUNKS
# ---------------------------------------------------------

def build_chunks(
    extracted_document,
    chunk_size=500,
    overlap=80
):
    """Create chunks while keeping source information."""

    all_chunks = []

    next_chunk_id = 1

    for page_info in extracted_document["pages"]:

        chunks = chunk_text(
            text=page_info["text"],
            source=extracted_document["source"],
            page_number=page_info["page"],
            start_chunk_id=next_chunk_id,
            chunk_size=chunk_size,
            overlap=overlap
        )

        all_chunks.extend(chunks)

        next_chunk_id += len(chunks)

    return all_chunks


# ---------------------------------------------------------
# SCORING
# ---------------------------------------------------------

def score_chunk(chunk, query):
    """Calculate simple relevance score."""

    query_words = meaningful_tokens(
        query
    )

    if not query_words:
        return 0, 0

    chunk_words = meaningful_tokens(
        chunk["text"]
    )

    matched_words = query_words.intersection(
        chunk_words
    )

    score = len(matched_words) * 5

    query_lower = clean_text(
        query
    ).lower()

    chunk_lower = chunk["text"].lower()

    if query_lower in chunk_lower:
        score += 20

    coverage = (
        len(matched_words) /
        len(query_words)
    )

    if coverage >= 0.75:

        score += 8

    elif coverage >= 0.50:

        score += 3

    return score, len(matched_words)


# ---------------------------------------------------------
# RETRIEVAL
# ---------------------------------------------------------

def retrieve_chunks(
    chunks,
    query,
    top_k=5
):
    """Retrieve most relevant chunks."""

    if not query or not query.strip():

        raise ValueError(
            "Query cannot be empty."
        )

    results = []

    for chunk in chunks:

        score, matched = score_chunk(
            chunk,
            query
        )

        if matched == 0:
            continue

        result = chunk.copy()

        result["score"] = score
        result["matched_terms"] = matched

        results.append(result)

    results.sort(
        key=lambda x: (
            x["score"],
            x["matched_terms"]
        ),
        reverse=True
    )

    return results[:top_k]


# ---------------------------------------------------------
# RAG CONTEXT
# ---------------------------------------------------------

def build_rag_context(
    retrieved_chunks
):
    """Build source-grounded context for Member 1."""

    if not retrieved_chunks:
        return ""

    context_parts = []

    for chunk in retrieved_chunks:

        page = chunk["page"]

        if page is None:
            page = "N/A"

        context_parts.append(

            f"[Source: {chunk['source']} | "
            f"Page: {page} | "
            f"Chunk: {chunk['chunk_id']} | "
            f"Score: {chunk['score']}]\n"
            f"{chunk['text']}"

        )

    return "\n\n".join(
        context_parts
    )


# ---------------------------------------------------------
# LLM READY INPUT
# ---------------------------------------------------------

def build_llm_ready_input(
    user_request,
    extracted_document,
    retrieved_chunks
):
    """Prepare clean input for Member 1's LLM layer."""

    context = build_rag_context(
        retrieved_chunks
    )

    return {

        "user_request":
            user_request,

        "source_file":
            extracted_document["source"],

        "source_type":
            extracted_document["file_type"],

        "context":
            context,

        "retrieved_chunks":
            retrieved_chunks,

        "grounding_available":
            bool(retrieved_chunks),

        "instructions":
            (
                "Use the provided source context "
                "as factual grounding. Do not invent "
                "source-specific facts. If the context "
                "is insufficient, clearly say so."
            )
    }


# ---------------------------------------------------------
# MAIN PROCESSING FUNCTION
# ---------------------------------------------------------

def process_document(
    file_path,
    user_request,
    top_k=5,
    chunk_size=500,
    overlap=80
):
    """
    MAIN FUNCTION FOR MEMBER 1.

    Example:

    result = process_document(
        "uploaded_file.pdf",
        "Create an executive summary."
    )
    """

    extracted = extract_document(
        file_path
    )

    chunks = build_chunks(
        extracted,
        chunk_size=chunk_size,
        overlap=overlap
    )

    retrieved = retrieve_chunks(
        chunks,
        user_request,
        top_k=top_k
    )

    llm_ready = build_llm_ready_input(
        user_request=user_request,
        extracted_document=extracted,
        retrieved_chunks=retrieved
    )

    return {

        "source":
            extracted["source"],

        "file_type":
            extracted["file_type"],

        "characters":
            extracted["characters"],

        "total_chunks":
            len(chunks),

        "retrieved_context":
            retrieved,

        "rag_context":
            llm_ready["context"],

        "llm_ready":
            llm_ready,

        "grounding_available":
            llm_ready[
                "grounding_available"
            ]
    }


# ---------------------------------------------------------
# DIRECT RUN
# ---------------------------------------------------------

if __name__ == "__main__":

    print(
        "Document Processor is ready."
    )

    print(
        "Use:"
    )

    print(
        "process_document(file_path, user_request)"
    )
