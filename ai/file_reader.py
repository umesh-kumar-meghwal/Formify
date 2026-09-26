import os


class FileReadError(Exception):
    """Raised when an uploaded file cannot be turned into text."""


def extract_text(file_path):
    """
    Extract plain text from an uploaded file.

    Supported:
    .txt, .md, .pdf, .docx, .jpg, .jpeg, .png
    """

    ext = os.path.splitext(file_path)[1].lower()

    try:

        # =========================
        # TXT / MARKDOWN
        # =========================
        if ext in (".txt", ".md"):

            with open(
                file_path,
                "r",
                encoding="utf-8",
                errors="ignore"
            ) as f:
                text = f.read()


        # =========================
        # PDF
        # =========================
        elif ext == ".pdf":

            from pypdf import PdfReader

            reader = PdfReader(file_path)

            text = "\n".join(
                (page.extract_text() or "")
                for page in reader.pages
            )


        # =========================
        # DOCX
        # =========================
        elif ext == ".docx":

            from docx import Document

            doc = Document(file_path)

            parts = [p.text for p in doc.paragraphs]

            # Read tables
            for table in doc.tables:

                for row in table.rows:

                    parts.append(
                        " | ".join(
                            cell.text
                            for cell in row.cells
                        )
                    )

            text = "\n".join(parts)


        # =========================
        # IMAGE OCR
        # =========================
        elif ext in (".jpg", ".jpeg", ".png"):

            import cv2
            import pytesseract

            pytesseract.pytesseract.tesseract_cmd = (
                r"C:\Program Files\Tesseract-OCR\tesseract.exe"
            )

            image = cv2.imread(file_path)

            if image is None:
                raise FileReadError(
                    "Could not open the image file."
                )

            # Upscale
            image = cv2.resize(
                image,
                None,
                fx=2,
                fy=2,
                interpolation=cv2.INTER_CUBIC
            )

            # Grayscale
            gray = cv2.cvtColor(
                image,
                cv2.COLOR_BGR2GRAY
            )

            # Improve contrast
            gray = cv2.threshold(
                gray,
                0,
                255,
                cv2.THRESH_BINARY + cv2.THRESH_OTSU
            )[1]

            # OCR
            text = pytesseract.image_to_string(
                gray,
                lang="eng",
                config="--psm 6"
            )


        # =========================
        # UNSUPPORTED FILE
        # =========================
        else:

            raise FileReadError(
                f"'{ext}' files cannot be read yet. "
                "Please upload a PDF, DOCX, TXT, MD, "
                "JPG, JPEG or PNG file."
            )


    except FileReadError:
        raise

    except Exception as e:

        raise FileReadError(
            f"Could not read the uploaded file: {e}"
        )


    # Remove unnecessary spaces
    text = text.strip()


    # No text found
    if not text:

        raise FileReadError(
            "No readable text was found in the file. "
            "The file may be scanned, image-only, "
            "blurry, or contain no readable text."
        )


    return text