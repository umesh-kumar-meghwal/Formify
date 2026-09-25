import os
import uuid

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


GENERATED_DIR = "generated_files"


def _label(key):
    return str(key).replace("_", " ").title()


def build_report_docx(data):
    """
    Works generically for both 'executive_summary' and 'advisory' JSON
    (from EXECUTIVE_SUMMARY_PROMPT / ADVISORY_PROMPT) since both are a
    title plus a mix of paragraph fields and list fields.
    """

    os.makedirs(GENERATED_DIR, exist_ok=True)

    doc = Document()

    title = data.get("title") or "Report"
    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for key, value in data.items():

        if key in ("output_type", "title"):
            continue

        if value in ("", None, []):
            continue

        doc.add_heading(_label(key), level=1)

        if isinstance(value, list):
            for item in value:
                doc.add_paragraph(str(item), style="List Bullet")
        else:
            doc.add_paragraph(str(value))

    filename = f"{uuid.uuid4().hex}.docx"
    output_path = os.path.join(GENERATED_DIR, filename)
    doc.save(output_path)

    return output_path